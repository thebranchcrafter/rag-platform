"""Local text extraction using pdfplumber, python-docx, etc."""

from typing import Dict, Any, List
from dataclasses import dataclass
import logging
import pdfplumber
import io
from pathlib import Path

logger = logging.getLogger(__name__)

# Optional imports
try:
    from docx import Document as DocxDocument
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False
    logger.warning("python-docx not available. DOCX files will not be supported.")

try:
    from bs4 import BeautifulSoup
    HTML_AVAILABLE = True
except ImportError:
    HTML_AVAILABLE = False
    logger.warning("beautifulsoup4 not available. HTML files will not be supported.")


@dataclass
class ExtractedContent:
    """Extracted content from document."""
    text: str
    tables: List[List[List[str]]]
    metadata: Dict[str, Any]
    confidence: float = 1.0
    pages_processed: int = 0


class LocalParserService:
    """Extract text using local parsers (pdfplumber, python-docx, etc.)."""
    
    def __init__(self):
        self.table_indicators = ['|', '\t', 'CH ', 'COEF', 'CUOTA', 'Tabla', 'CHALET']
    
    async def extract(
        self,
        file_path: str,
        file_type: str
    ) -> ExtractedContent:
        """Extract text from document using appropriate local parser."""
        logger.info(f"[Local Parser] Starting extraction for {file_type} file: {file_path}")
        logger.debug(f"[Local Parser] Available parsers - PDF: pdfplumber, DOCX: {DOCX_AVAILABLE}, HTML: {HTML_AVAILABLE}")
        
        if file_type == "pdf":
            logger.debug("[Local Parser] Using pdfplumber for PDF extraction")
            return await self._extract_pdf(file_path)
        elif file_type == "docx":
            if not DOCX_AVAILABLE:
                raise ValueError("python-docx not available. Install with: pip install python-docx")
            logger.debug("[Local Parser] Using python-docx for DOCX extraction")
            return await self._extract_docx(file_path)
        elif file_type == "html":
            if not HTML_AVAILABLE:
                raise ValueError("beautifulsoup4 not available. Install with: pip install beautifulsoup4")
            logger.debug("[Local Parser] Using BeautifulSoup for HTML extraction")
            return await self._extract_html(file_path)
        else:
            raise ValueError(f"Unsupported file type for local parser: {file_type}")
    
    async def _extract_pdf(self, file_path: str) -> ExtractedContent:
        """Extract text from PDF using pdfplumber."""
        logger.debug(f"[Local Parser/PDF] Opening PDF file: {file_path}")
        text_parts = []
        all_tables = []
        pages_processed = 0
        tables_found = 0
        
        try:
            with open(file_path, 'rb') as f:
                pdf_file = io.BytesIO(f.read())
            
            with pdfplumber.open(pdf_file) as pdf:
                total_pages = len(pdf.pages)
                logger.info(f"[Local Parser/PDF] Processing {total_pages} pages with pdfplumber")
                
                for page_num, page in enumerate(pdf.pages, 1):
                    # Extract regular text
                    page_text = page.extract_text()
                    
                    if page_text:
                        text_parts.append(page_text)
                        pages_processed += 1
                        logger.debug(f"[Local Parser/PDF] Page {page_num}: Extracted {len(page_text)} characters")
                        
                        # Check if page likely contains tables
                        has_table_indicators = any(
                            keyword in page_text for keyword in self.table_indicators
                        )
                        
                        if has_table_indicators:
                            logger.debug(f"[Local Parser/PDF] Page {page_num}: Table indicators detected, extracting tables...")
                            # Extract tables
                            tables = page.extract_tables()
                            if tables:
                                all_tables.extend(tables)
                                tables_found += len(tables)
                                logger.debug(f"[Local Parser/PDF] Page {page_num}: Found {len(tables)} table(s)")
                                # Format tables as text
                                for table_num, table in enumerate(tables, 1):
                                    if table and len(table) > 0:
                                        table_text = self._format_table(table)
                                        if table_text:
                                            text_parts.append(f"\n\n[Tabla {table_num}]\n{table_text}\n")
            
            full_text = "\n".join(text_parts).strip()
            
            logger.info(f"[Local Parser/PDF] Extraction complete:")
            logger.info(f"  - Pages processed: {pages_processed}/{total_pages}")
            logger.info(f"  - Total text length: {len(full_text)} characters")
            logger.info(f"  - Tables found: {tables_found}")
            logger.info(f"  - Confidence: {'1.0' if full_text else '0.0'}")
            
            return ExtractedContent(
                text=full_text,
                tables=all_tables,
                metadata={
                    "extraction_method": "pdfplumber",
                    "pages_processed": pages_processed,
                    "total_pages": total_pages,
                    "tables_found": tables_found
                },
                confidence=1.0 if full_text else 0.0,
                pages_processed=pages_processed
            )
        
        except Exception as e:
            logger.error(f"[Local Parser/PDF] Error extracting PDF: {str(e)}", exc_info=True)
            raise
    
    async def _extract_docx(self, file_path: str) -> ExtractedContent:
        """Extract text from DOCX using python-docx."""
        if not DOCX_AVAILABLE:
            raise ValueError("python-docx not available. Install with: pip install python-docx")
        
        try:
            doc = DocxDocument(file_path)
            
            text_parts = []
            tables = []
            
            # Extract paragraphs
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_parts.append(paragraph.text)
            
            # Extract tables
            for table in doc.tables:
                table_data = []
                for row in table.rows:
                    row_data = [cell.text.strip() for cell in row.cells]
                    table_data.append(row_data)
                tables.append(table_data)
                
                # Format table as text
                table_text = self._format_table(table_data)
                if table_text:
                    text_parts.append(f"\n\n[Tabla]\n{table_text}\n")
            
            full_text = "\n".join(text_parts).strip()
            
            return ExtractedContent(
                text=full_text,
                tables=tables,
                metadata={
                    "extraction_method": "python-docx",
                    "paragraphs": len(doc.paragraphs)
                },
                confidence=1.0,
                pages_processed=1  # DOCX doesn't have pages
            )
        
        except Exception as e:
            logger.error(f"Error extracting DOCX: {str(e)}")
            raise
    
    async def _extract_html(self, file_path: str) -> ExtractedContent:
        """Extract text from HTML using BeautifulSoup."""
        if not HTML_AVAILABLE:
            raise ValueError("beautifulsoup4 not available. Install with: pip install beautifulsoup4")
        
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                html_content = f.read()
            
            soup = BeautifulSoup(html_content, 'html.parser')
            
            # Remove script and style elements
            for script in soup(["script", "style"]):
                script.decompose()
            
            # Get text
            text = soup.get_text(separator='\n', strip=True)
            
            # Extract tables
            tables = []
            for table in soup.find_all('table'):
                table_data = []
                for row in table.find_all('tr'):
                    row_data = [cell.get_text(strip=True) for cell in row.find_all(['td', 'th'])]
                    if row_data:
                        table_data.append(row_data)
                if table_data:
                    tables.append(table_data)
            
            return ExtractedContent(
                text=text,
                tables=tables,
                metadata={
                    "extraction_method": "beautifulsoup4"
                },
                confidence=1.0,
                pages_processed=1
            )
        
        except Exception as e:
            logger.error(f"Error extracting HTML: {str(e)}")
            raise
    
    def _format_table(self, table: List[List[str]]) -> str:
        """Format table as readable text."""
        if not table or len(table) == 0:
            return ""
        
        # Filter out empty rows
        table = [row for row in table if any(cell and str(cell).strip() for cell in row)]
        
        if not table:
            return ""
        
        # Normalize table
        max_cols = max(len(row) for row in table) if table else 0
        normalized_table = []
        for row in table:
            normalized_row = [str(cell).strip() if cell else "" for cell in row]
            while len(normalized_row) < max_cols:
                normalized_row.append("")
            normalized_table.append(normalized_row)
        
        # Format as pipe-separated
        formatted_lines = []
        for row in normalized_table:
            if any(cell.strip() for cell in row):
                row_text = " | ".join(cell if cell else "" for cell in row)
                formatted_lines.append(row_text)
        
        return "\n".join(formatted_lines)
